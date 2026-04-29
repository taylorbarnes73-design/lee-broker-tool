#!/usr/bin/env python3
"""Sample the `photo` column from both sheets to understand its format."""
from openpyxl import load_workbook

path = '/home/ubuntu/upload/LeeIrvine_Metropolis_CompData.xlsx'
wb = load_workbook(path, read_only=True, data_only=True)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    # find the photo column index
    photo_idx = None
    header = None
    for row in ws.iter_rows(values_only=True, min_row=1, max_row=1):
        header = row
        for i, h in enumerate(row):
            if h == 'photo':
                photo_idx = i
                break
        break
    if photo_idx is None:
        print(f'{sheet_name}: no photo column')
        continue

    print(f'--- {sheet_name} (photo col idx={photo_idx}) ---')
    non_empty = 0
    samples = []
    total = 0
    for i, row in enumerate(ws.iter_rows(values_only=True, min_row=2)):
        total += 1
        if photo_idx < len(row):
            val = row[photo_idx]
            if val and str(val).strip() and str(val).strip().upper() != 'NULL':
                non_empty += 1
                if len(samples) < 15:
                    samples.append(val)
    print(f'  total rows: {total}, non-empty photos: {non_empty}')
    for s in samples:
        print(f'  sample: {s!r}')
